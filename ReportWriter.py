from docx import Document
from docx.shared import Inches
from docx.shared import Pt
import pandas as pd
from docx.enum.table import WD_ALIGN_VERTICAL
from pathlib import Path

from docx.shared import Inches
import numpy

# creates a sqa analysis document for each subject and level
# with hisotric grades and components

class ReportWriter:
    def __init__(self,subject,level,year,years,grades,school_comp,national_comp):
        self.years = years
        self.doc = Document()
        if(level==75):
            self.doc.add_heading(subject+": National 5 "+str(year), level=1)
        else:
            self.doc.add_heading(subject+": Higher", level=1)
        self.doc.add_heading('Overview', level=2)
        self.doc.add_paragraph("<insert discussion>")
        
        self.doc.add_heading('School Vs National Comparisons', level=2)
        self.doc.add_heading('Grades', level=3)
        table = self.doc.add_table(rows=0, cols=2,style='Table Grid')
        fname = self.imageNameFinder(subject,level,"comp")
        subject_grades = grades[grades["Year"]==year]
        subject_grades = subject_grades["Num"].values
        self.gradeComparison(table,fname,subject_grades,school_comp[year],national_comp[year])
        self.doc.add_heading('Component Marks', level=3)
        self.componentComparison(school_comp[year],national_comp[year])
        
        self.doc.add_page_break()
        self.doc.add_heading('Historic Comparisons', level=2)
        
        num_entries = {}
        
        for year in years:
            subject_grades = grades[grades["Year"]==year]
            subject_grades = subject_grades["Num"].values
            num_entries[year] = subject_grades.sum()

        self.historicGrades(subject,level,num_entries)
        self.doc.add_heading('Component Marks (mean)', level=3)
        self.historicComp(school_comp,national_comp)
        self.fname = "Output/Subject Overviews/"
        if(level==75):
            self.fname = self.fname +subject+"_N5"
        else:
            self.fname = self.fname +subject+"_H"
        self.fname += ".docx"
        self.save()
    
    def createCompTable(self,data,name,years,num_comp,component_names):
        comp_table = self.doc.add_table(rows=num_comp+1, cols=4,style='Table Grid')
        cells = comp_table.rows[0].cells
        cells[0].text = name
        for i in range(1,4):
            cells[i].text = str(years[i-1])
        cells = comp_table.columns[0].cells
        for i in range(1,num_comp+1):
            cells[i].text = str(component_names[i])

        for i in range(1,4):
            year = years[i-1]
        
            cells = comp_table.columns[i].cells
            for j in range(1,num_comp+1):
                s_mean = data[year]["mean"].get(j,"N/A")
                cells[j].text = str(s_mean)



    def historicComp(self,school,national):
        num_comp = max(list(school[self.years[0]]['min'].keys()))
        comp_names = school[self.years[0]]["Title"]
        self.doc.add_heading('School', level=3)
        self.createCompTable(school,"Components",self.years,num_comp,comp_names)
        self.doc.add_heading('National', level=3)
        self.createCompTable(national,"Components",self.years,num_comp,comp_names)
        self.doc.add_heading('Difference(School-National)', level=3)
        difference = {}
        for year in self.years:
            s_data = school[year]['mean']
            n_data = national[year]['mean']
            diff = {}
            for comp in s_data:
                temp = s_data[comp]-n_data[comp]
                diff[comp] = round(temp,1)
            
            difference[year] = {"mean":diff}
        
        self.createCompTable(difference,"Components",self.years,num_comp,comp_names)
        
        pass
    

    def historicGrades(self,subject,level,num_entries):
        self.doc.add_heading('School Grades', level=3)
        fname = self.imageNameFinder(subject,level,"hist")
        fname_school = fname+"school.png"
        fname_national = fname+"national.png"
        
        table = self.doc.add_table(rows=0, cols=2,style='Table Grid')
        row_cells = table.add_row().cells
        paragraph = row_cells[0].paragraphs[0]
        paragraph.style=None
        run = paragraph.add_run(style=None)
        run.add_picture(fname_school,width=Inches(4.5))
        row_cells[1].paragraphs[0].style=None
        row_cells[1]._element.clear_content()
        p =row_cells[1].add_paragraph('\n\n\n Num Entries')
        p.alignment = 1
        table = row_cells[1].add_table(rows=3,cols=2)
        table.style='Table Grid'
        
        pos = 0
        for year in self.years:
            table.rows[pos].cells[0].text = str(year)
            table.columns[1].cells[pos].text = str(num_entries[year])
            pos +=1
        
        self.doc.add_heading('National Grades', level=3)
        self.doc.add_picture(fname_national,width=Inches(4.5))

        
        
        
        

    def imageNameFinder(self,subject,level,type):
        root = "Output/reportgraphs/"
        fname = ""
        if(type=="comp"):
            fname = root+"/comp/"+(subject+"_"+str(level)+".png")
        elif(type=="hist"):
            fname = root+"/historic/"+(subject+"_"+str(level)+"_")
        else:
            print(type," type not known")
            exit()
        return fname

    def save(self):
        self.doc.save(self.fname)

    def gradeComparison(self,table,graph_fname,gradecounts,school_comp,national_comp):
        row_cells = table.add_row().cells
        paragraph = row_cells[0].paragraphs[0]
        paragraph.style=None
        run = paragraph.add_run(style=None)
        run.add_picture(graph_fname,width=Inches(4.5))
        row_cells[1].paragraphs[0].style=None
        row_cells[1]._element.clear_content()
        p =row_cells[1].add_paragraph('\n\n\n School Grade Counts')
        p.alignment = 1
        table = row_cells[1].add_table(rows=5,cols=2)
        table.style='Table Grid'
        header = table.rows[0].cells
        header[0].text="Grade"
        header[1].text="School"
        table.rows[0].cells[0].text = "A"
        table.rows[1].cells[0].text = "B"
        table.rows[2].cells[0].text = "C"
        table.rows[3].cells[0].text = "D"
        table.rows[4].cells[0].text = "NA"
        for i in range(0,5):
            table.columns[1].cells[i].text = str(gradecounts[i])

    def writeCol(self,table,labels,positions,num_comp,data):
        for i in range(len(labels)):
            label = labels[i]
            pos = positions[i]
            temp = data[label]
            cells = table.columns[pos-1].cells
            for j in range(1,num_comp+1):
                cells[j].text = str(temp[j])

    def componentComparison(self,school,national):
        num_comp = max(list(school['min'].keys()))
        comp_table = self.doc.add_table(rows=num_comp+1, cols=6,style='Table Grid')
        cells = comp_table.rows[0].cells
        cells[1].text = "School Mean"
        cells[2].text = "National Mean"
        cells[3].text = "Difference"
        cells[4].text = "School Max"
        cells[5].text = "National Max"
        
        temp = {}
        for i in range(1,num_comp+1):
            temp[i] = round(school["mean"][i] -national["mean"][i],2)
        diff = {"diff":temp}

        labels = ["Title","mean","max"]
        positions = [1,2,5]
        self.writeCol(comp_table,labels,positions,num_comp,school)
        
        labels = ["mean","maximum"]
        positions = [3,6]
        self.writeCol(comp_table,labels,positions,num_comp,national)

        labels = ["diff"]
        positions = [4]
        self.writeCol(comp_table,labels,positions,num_comp,diff)
        

if __name__=="__main__":
    ReportWriter("Biology",75,2024,[3,4,5,6,1],[],[])

